"""Lossless workbook transport, with visible columns and a strict remainder field."""
from __future__ import annotations

import copy
import io
import json
import uuid
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from ..authoring_service import AuthoringError, require
from .exchange import archive_entries, FORMAT

SHEETS={'cards':'Ziele','sentences':'Saetze','card_sentences':'Luecken',
        'sentence_tokens':'Tokens','lemmas':'Lemmas','senses':'Bedeutungen',
        'dictionary_forms':'Formglossen','decks':'Stapel','provenance':'Herkunft',
        'operations':'Operationen'}
VISIBLE=('ref','card','sentence','deck','form','translation_de','text','sense','lemma',
         'sense_key','definition_de','gloss_de','pos','idx','surface','start_pos','end_pos',
         'gap_start','gap_end','context_gloss_de','title_de','slug','op_id','kind','target_ref')


def _cell(sheet,row,column,value):
    if value is not None and not isinstance(value,(str,int,float,bool)):
        value=json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':'))
    require(not isinstance(value,str) or len(value)<=32767,'XLSX_CELL_LIMIT',f'{sheet.title}, Zeile {row}, Spalte {column}: Eine Tabellenzelle überschreitet 32.767 Zeichen. JSON/ZIP verwenden.')
    cell=sheet.cell(row,column,value)
    if isinstance(value,str):cell.data_type='s'


def _write_rows(book,name,rows,columns):
    sheet=book.create_sheet(name)
    for j,k in enumerate(columns,1):_cell(sheet,1,j,k)
    for i,row in enumerate(rows,2):
        for j,k in enumerate(columns,1):_cell(sheet,i,j,row.get(k))
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='273B4A')
    for column in sheet.columns:
        letter=column[0].column_letter
        sheet.column_dimensions[letter].width=36 if column[0].value not in ('idx','start_pos','end_pos') else 14
    return sheet


def write_workbook(delivery):
    book=Workbook();book.remove(book.active)
    metadata={k:copy.deepcopy(v) for k,v in delivery.items() if k not in ('add','operations','provenance')}
    metadata['xlsx_mode']='full';metadata['xlsx_version']=1
    metadata['xlsx_add_tables']=list(delivery.get('add',{}))
    _write_rows(book,'Meta',[{'key':k,'value_json':json.dumps(v,ensure_ascii=False,sort_keys=True)} for k,v in metadata.items()],['key','value_json'])
    grouped={name:[] for name in [*SHEETS.values(),'Bindungen','Gruppen','Alternativen','Records']}
    for table,records in {**delivery.get('add',{}),'operations':delivery.get('operations',[]),'provenance':delivery.get('provenance',[])}.items():
        for index,record in enumerate(records):
            rest=copy.deepcopy(record)
            if table=='cards' and 'learning' in rest:
                grouped['Gruppen'].append({'_table':table,'_index':index,'payload_json':json.dumps({'learning':rest.pop('learning')},ensure_ascii=False)})
            if table=='card_sentences' and 'valid_alternatives' in rest:
                grouped['Alternativen'].append({'_table':table,'_index':index,'payload_json':json.dumps({'valid_alternatives':rest.pop('valid_alternatives')},ensure_ascii=False)})
            row={'_table':table,'_index':index}
            for k in VISIBLE:
                if k in rest:row[k]=rest.pop(k)
            # Track explicit null vs absent in the remainder, not a guessed value.
            row['payload_json']=json.dumps({'fields':rest,'present':[k for k in VISIBLE if k in record]},ensure_ascii=False,sort_keys=True)
            name=SHEETS.get(table,'Bindungen' if table in ('deck_cards','deck_words','word_aliases','card_lexical_targets','deck_practice_targets') else 'Records')
            grouped[name].append(row)
    for name,rows in grouped.items():
        cols=['_table','_index']+[k for k in VISIBLE if any(k in r for r in rows)]+['payload_json']
        _write_rows(book,name,rows,cols)
    _write_rows(book,'Anleitung',[{'Hinweis':line} for line in FORMAT.splitlines() if line],['Hinweis'])
    buffer=io.BytesIO();book.save(buffer);return buffer.getvalue()


def _rows(sheet):
    values=list(sheet.iter_rows())
    if not values:return []
    headers=[c.value for c in values[0]]
    require(all(isinstance(h,str) for h in headers) and len(set(headers))==len(headers),'XLSX_COLUMNS',sheet.title+': ungültige oder doppelte Spalten.')
    result=[]
    for number,cells in enumerate(values[1:],2):
        if all(c.value is None for c in cells):continue
        for c in cells:
            if c.data_type=='f':
                raise AuthoringError('XLSX_FORMULA',f'{sheet.title}!{c.coordinate}: Formeln sind nicht zulässig.',field=c.coordinate,record_ref=sheet.title)
        result.append((number,{h:c.value for h,c in zip(headers,cells)}))
    return result


def read_workbook(raw,lang):
    archive_entries(raw,xlsx=True)
    try:book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    except Exception as e:raise AuthoringError('FORMAT_INVALID','XLSX kann nicht gelesen werden: '+str(e)) from e
    try:
        require('Meta' in book and 'Ziele' in book,'XLSX_SHEET','Meta und Ziele sind erforderlich.')
        meta={}
        for row,record in _rows(book['Meta']):
            key=record.get('key');require(key and key not in meta,'XLSX_META','Doppelte oder leere Metaangabe.')
            try:meta[key]=json.loads(record.get('value_json') or 'null')
            except ValueError:raise AuthoringError('XLSX_META',f'Meta Zeile {row}: ungültiger Wert.') from None
        mode=meta.pop('xlsx_mode','draft');version=meta.pop('xlsx_version',1)
        add_tables=meta.pop('xlsx_add_tables',[])
        require(version==1 and mode in ('full','draft'),'FORMAT_UNSUPPORTED','Unbekannte XLSX-Version oder Modus.')
        if mode=='draft':return _simple(book,meta,lang)
        expected={'Meta','Anleitung','Bindungen','Gruppen','Alternativen','Records',*SHEETS.values()}
        require(set(book.sheetnames)==expected,'XLSX_SHEET','Vollständige Lieferung: Detailblätter fehlen oder sind unbekannt.')
        records={};parts=[]
        for sheet in book:
            if sheet.title in ('Meta','Anleitung'):continue
            for number,row in _rows(sheet):
                require(set(row)<={'_table','_index','payload_json',*VISIBLE},'XLSX_COLUMNS',sheet.title+': unbekannte Spalte.')
                table,index=row.get('_table'),row.get('_index')
                require(isinstance(table,str) and type(index) is int and index>=0,'XLSX_ROW',f'{sheet.title} Zeile {number}: Zeilenreferenz fehlt.')
                try:payload=json.loads(row.get('payload_json') or '{}')
                except ValueError:raise AuthoringError('XLSX_ROW',f'{sheet.title} Zeile {number}: ungültige Restfelder.') from None
                if sheet.title in ('Gruppen','Alternativen'):
                    parts.append((table,index,payload));continue
                require(isinstance(payload,dict) and set(payload)=={'fields','present'},'XLSX_ROW','Ungültiges Restfeldformat.')
                record=payload['fields'];present=payload['present']
                require(isinstance(record,dict) and not set(record)&set(VISIBLE),'XLSX_OVERLAP','Sichtbare Spalte und Restfeld widersprechen sich.')
                require(isinstance(present,list) and set(present)<=set(VISIBLE),'XLSX_ROW','Ungültige Feldliste.')
                for k in VISIBLE:
                    if k in present or row.get(k) is not None:record[k]=row.get(k)
                require((table,index) not in records,'XLSX_ROW','Doppelte Zeilenreferenz.')
                records[table,index]=record
        for table,index,payload in parts:
            require((table,index) in records and isinstance(payload,dict) and not set(records[table,index])&set(payload),'XLSX_OVERLAP','Doppelte oder verwaiste Detailbindung.')
            records[table,index].update(payload)
        meta.update(add={t:[] for t in add_tables},operations=[],provenance=[])
        for table in sorted({t for t,_ in records}):
            indexes=sorted(i for t,i in records if t==table)
            require(indexes==list(range(len(indexes))),'XLSX_ROW','Zeilenindex ist nicht vollständig: '+table)
            rows=[records[table,i] for i in indexes]
            if table in ('operations','provenance'):meta[table]=rows
            else:meta['add'][table]=rows
        return meta
    finally:book.close()


def _simple(book,meta,lang):
    allowed={'ref','form','translation_de','sentence','sentence_translation_de','sense','gap_start','gap_end'}
    cards=[];sentences=[];links=[];locations={}
    delivery_id=meta.get('delivery_id') or str(uuid.uuid4())
    for number,row in _rows(book['Ziele']):
        require(set(row)<=allowed,'XLSX_COLUMNS','Einfacher Modus: unbekannte Zielspalte.')
        ref=row.get('ref') or delivery_id+':row:'+str(number);sn='sentence:'+ref
        cards.append({'ref':ref,'form':row.get('form'),'translation_de':row.get('translation_de'),'sense':row.get('sense')})
        sentences.append({'ref':sn,'text':row.get('sentence'),'translation_de':row.get('sentence_translation_de')})
        links.append({'card':ref,'sentence':sn,'position':1,'gap_start':row.get('gap_start'),'gap_end':row.get('gap_end')})
        locations[ref]={'sheet':'Ziele','row':number}
    return {'format':'sprachapp.editorial-content','format_version':4,'delivery_id':delivery_id,
            'lang':meta.get('lang',lang),'translation_lang':'de','purpose':'append',
            'base':meta.get('base',{}),'context_sha256':meta.get('context_sha256'),
            'expected_target_refs':[r['ref'] for r in cards],
            'add':{'cards':cards,'sentences':sentences,'card_sentences':links},'operations':[],
            'provenance':[],'submitted_reviews':[],'extensions':{},'locations':locations,'completeness':'draft'}
