import io
import json
import pytest
from openpyxl import Workbook,load_workbook
from sprachpipe.authoring_service import AuthoringError,empty_pack
from sprachpipe.editor.exchange import template_delivery
from sprachpipe.editor.xlsx import write_workbook,read_workbook


def test_full_roundtrip_nested_bindings_and_formula_like_strings():
    d,_=template_delivery(empty_pack('en','English'),2)
    d['add']['cards'][0].update(form='=not a formula',learning={'group_id':'g','related':[]},extra={'keep':[1,None,False]})
    d['add']['card_sentences']=[{'card':d['expected_target_refs'][0],'sentence':'s','gap_start':0,'gap_end':3,'valid_alternatives':[]}]
    d['add']['sentence_tokens']=[{'sentence':'s','idx':0,'surface':'x','context_gloss_de':'@literal','context_snapshot':{'definition_de':'test'}}]
    d['operations']=[{'op_id':'x','kind':'set_solution_hint','expect':{'x':'y'},'payload':{'translation_de':'-literal'}}]
    d['provenance']=[{'ref':'p','kind':'manual','notes':'keep'}]
    assert read_workbook(write_workbook(d),'en')==d


def test_formula_and_overlap_are_rejected():
    d,_=template_delivery(empty_pack('en','English'),1)
    book=load_workbook(io.BytesIO(write_workbook(d)))
    book['Ziele']['C2']='=1+1'
    out=io.BytesIO();book.save(out)
    with pytest.raises(AuthoringError,match='Formeln'):read_workbook(out.getvalue(),'en')


def test_simple_workbook_keeps_missing_annotations():
    book=Workbook();book.active.title='Meta';book['Meta'].append(['key','value_json']);book['Meta'].append(['xlsx_mode','"draft"'])
    sheet=book.create_sheet('Ziele');sheet.append(['form','translation_de','sentence','sentence_translation_de'])
    sheet.append(['cut','schneiden','They cut paper.','Sie schneiden Papier.'])
    out=io.BytesIO();book.save(out);d=read_workbook(out.getvalue(),'en')
    assert d['add']['cards'][0]['sense'] is None
    assert 'sentence_tokens' not in d['add']
    assert 'valid_alternatives' not in d['add']['card_sentences'][0]
    assert d['locations'][d['expected_target_refs'][0]]=={'sheet':'Ziele','row':2}


def test_long_cell_is_never_silently_truncated():
    d,_=template_delivery(empty_pack('en','English'),1);d['add']['cards'][0]['form']='x'*32768
    with pytest.raises(AuthoringError,match='32.767'):write_workbook(d)
